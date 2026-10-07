"""Extract plain text from an uploaded resume (PDF, DOCX or TXT)."""

import io
import logging

log = logging.getLogger(__name__)

MAX_BYTES = 5 * 1024 * 1024
MAX_PAGES = 10


class ExtractError(Exception):
    pass


def extract_text(filename: str, content: bytes) -> str:
    if len(content) > MAX_BYTES:
        raise ExtractError("File is too large (max 5 MB).")
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or content.startswith(b"%PDF"):
            return _pdf(content)
        if name.endswith(".docx") or content.startswith(b"PK"):
            return _docx(content)
        if name.endswith((".txt", ".md")):
            return content.decode("utf-8", errors="replace")
    except ExtractError:
        raise
    except Exception as e:  # malformed files raise a zoo of parser errors
        log.info("Could not parse upload %r: %s", filename, e)
        raise ExtractError("We couldn't read that file. Try exporting it as PDF or DOCX.") from e
    raise ExtractError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def _pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ExtractError("That PDF is password-protected.")
    pages = reader.pages[:MAX_PAGES]
    return "\n".join((p.extract_text() or "") for p in pages)


def _docx(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)
